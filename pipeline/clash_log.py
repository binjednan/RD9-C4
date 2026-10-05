# -*- coding: utf-8 -*-
"""clash register: stable ids, statuses coming from the site log, change history.

  pipeline/data/clash_log.json   {"version":1,"entries":{<id>:{status,res,note,date,by,ev}},"history":[{ts,id,field,old,new,by,src}]}
  python3 pipeline/clash_log.py set <clash-id|ea|eb> <status> "<note>" [--res "<how it was solved>"] [--by NAME] [--date YYYY-MM-DD] [--src FILE]
  python3 pipeline/clash_log.py import <file.csv|xlsx>   # columns (any order, Arabic or English): id | element A | element B | status | resolution | note | date | by
  then  python3 pipeline/post_model.py && python3 src/build.py     # merges the log into model.json (clashes[].st/res/note/date/by) and rewrites docs/CLASH_LOG.md
Every change appends an entry to 'history' (old -> new) so the register keeps a full audit trail."""
import json, os, sys, zlib, datetime, csv, re
HERE = os.path.dirname(os.path.abspath(__file__))
LOG = os.path.join(HERE, "data", "clash_log.json")
STATUS_AR = {"open": "مفتوح — لم يُفحص", "checking": "قيد الفحص", "resolved": "حُلّ في التنفيذ", "accepted": "مقبول كما هو (لا مشكلة)", "false": "ليس تعارضًا (إيجابي كاذب)", "design": "يحتاج قرارًا تصميميًا"}
ALIASES = {"open": "open", "مفتوح": "open", "new": "open", "checking": "checking", "قيد الفحص": "checking", "in progress": "checking", "resolved": "resolved", "closed": "resolved", "حل": "resolved", "تم الحل": "resolved", "محلول": "resolved", "حُلّ": "resolved",
           "accepted": "accepted", "مقبول": "accepted", "false": "false", "false positive": "false", "ليس تعارضا": "false", "ليس تعارضًا": "false", "design": "design", "قرار": "design"}

def clash_id(ea, eb):
    return "CL-%08X" % (zlib.crc32(f"{ea}|{eb}".encode()) & 0xffffffff)

def load():
    if os.path.exists(LOG):
        return json.load(open(LOG, encoding="utf-8"))
    return {"version": 1, "entries": {}, "history": []}

def save(d):
    json.dump(d, open(LOG, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

def norm_status(s):
    s = (s or "").strip().lower()
    return ALIASES.get(s) or ALIASES.get(s.replace("ً", "").replace("ّ", "")) or ("open" if not s else None)

def set_entry(d, cid, status=None, note=None, res=None, by=None, date=None, src=None, ev=None):
    e = d["entries"].setdefault(cid, {})
    ts = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    for k, v in (("status", status), ("note", note), ("res", res), ("by", by), ("date", date), ("ev", ev)):
        if v is None: continue
        old = e.get(k)
        if old != v:
            d["history"].append({"ts": ts, "id": cid, "field": k, "old": old, "new": v, "by": by or "", "src": src or ""})
            e[k] = v
    return e

def merge_into(M, els):
    """called by post_model.py: stamps ids and statuses on M['clashes'], returns summary"""
    d = load()
    ent = d["entries"]; cnt = {}
    for c in M["clashes"]:
        ea, eb = els[c["a"]]["id"], els[c["b"]]["id"]
        c["id"] = clash_id(ea, eb); c["ea"] = ea; c["eb"] = eb
        e = ent.get(c["id"]) or ent.get(f"{ea}|{eb}") or ent.get(f"{eb}|{ea}") or {}
        c["st"] = e.get("status", "open")
        for k in ("res", "note", "date", "by", "ev"):
            if e.get(k): c[k if k != "ev" else "ev"] = e[k]
        cnt[c["st"]] = cnt.get(c["st"], 0) + 1
    known = {c["id"] for c in M["clashes"]} | {f"{c['ea']}|{c['eb']}" for c in M["clashes"]} | {f"{c['eb']}|{c['ea']}" for c in M["clashes"]}
    orphans = [k for k in ent if k not in known]
    M["clashLog"] = {"version": d.get("version", 1), "statuses": STATUS_AR, "counts": cnt, "history": d["history"][-200:], "orphans": orphans,
                     "updated": (d["history"][-1]["ts"] if d["history"] else None), "source": d.get("source")}
    return M["clashLog"]

def write_doc(M):
    L = M["clashLog"]; out = ["# سجل التعارضات — الحالة والتغييرات", "",
         "يُولَّد هذا الملف آليًا من `pipeline/data/clash_log.json` عند تشغيل `pipeline/post_model.py`. لا يُعدَّل يدويًا.", "",
         f"إجمالي التعارضات المرصودة: **{len(M['clashes'])}**", "", "| الحالة | العدد |", "|---|---|"]
    for k, v in STATUS_AR.items(): out.append(f"| {v} | {L['counts'].get(k, 0)} |")
    out += ["", "## آخر التغييرات المسجّلة", ""]
    if not L["history"]: out.append("_لا توجد تغييرات مسجّلة بعد — بانتظار ملف السجلات من الموقع._")
    else:
        out += ["| التاريخ | التعارض | الحقل | من | إلى | بواسطة | المصدر |", "|---|---|---|---|---|---|---|"]
        for h in reversed(L["history"][-60:]): out.append(f"| {h['ts']} | {h['id']} | {h['field']} | {h.get('old') or '—'} | {h['new']} | {h.get('by') or ''} | {h.get('src') or ''} |")
    if L["orphans"]: out += ["", "## مدخلات في السجل لا تطابق أي تعارض حالي (تحتاج مراجعة)", ""] + [f"- `{k}`" for k in L["orphans"]]
    open(os.path.join(os.path.dirname(HERE), "docs", "CLASH_LOG.md"), "w", encoding="utf-8").write("\n".join(out) + "\n")

def _cli():
    a = sys.argv[1:]
    if not a: print(__doc__); return
    d = load()
    if a[0] == "set":
        cid, st, note = a[1], norm_status(a[2]), (a[3] if len(a) > 3 and not a[3].startswith("--") else None)
        kw = {}; i = 3 if note is None else 4
        while i < len(a) - 1:
            if a[i].startswith("--"): kw[a[i][2:]] = a[i + 1]; i += 2
            else: i += 1
        if st is None: print("unknown status; use:", ", ".join(STATUS_AR)); return
        set_entry(d, cid, st, note, kw.get("res"), kw.get("by"), kw.get("date"), kw.get("src")); save(d); print("ok")
    elif a[0] == "import":
        path = a[1]; rows = []
        if path.lower().endswith(".csv"):
            rows = list(csv.DictReader(open(path, encoding="utf-8-sig")))
        else:
            import openpyxl
            ws = openpyxl.load_workbook(path, data_only=True).active; hdr = [str(c.value or "").strip() for c in ws[1]]
            rows = [dict(zip(hdr, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
        def pick(r, *names):
            for k, v in r.items():
                kk = (k or "").strip().lower()
                if any(n in kk for n in names) and v not in (None, ""): return str(v).strip()
        n = 0
        for r in rows:
            cid = pick(r, "id", "رقم", "معرف")
            if not cid: continue
            st = norm_status(pick(r, "status", "الحالة", "حالة"))
            set_entry(d, cid, st, pick(r, "note", "ملاحظ"), pick(r, "resol", "الحل", "حل"), pick(r, "by", "بواسطة", "المهندس"), pick(r, "date", "تاريخ"), os.path.basename(path)); n += 1
        d["source"] = os.path.basename(path); save(d); print("imported", n, "rows")

if __name__ == "__main__":
    _cli()
